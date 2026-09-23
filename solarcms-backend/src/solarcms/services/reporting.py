"""Report rendering: query the aggregate tiers, produce XLSX and PDF.

BACKEND_SPEC §13 Phase 9, MASTER §6.6.

Three rules shape everything here:

* **Reports render from aggregate tiers, never raw `readings`.** The tier is
  chosen by `domain/tiering.select_tier`, like every other reader.
* **I-11: Financial Reports are computed from ABT Meter Readings, never MFM.**
  The ABT Meter is the sealed, revenue-grade settlement instrument; an MFM is
  operational monitoring with no commercial standing. A Financial Report with no
  ABT Meter available *fails* rather than falling back — producing an invoice
  figure from the wrong instrument is worse than producing none.
* **Energy is the same figure the KPI screen shows.** It comes from
  `domain/counters`: each meter integrated on its own, one Device Type per Plant
  by `PLANT_ENERGY_COUNTER_PRECEDENCE`, and every counter step that went
  backwards or jumped further than the Plant could produce refused and *listed*.
  ⚠ This used to sum `max - min` over every meter of the Client, which counted a
  Plant with a settlement meter and a check meter twice and turned every counter
  reset into gigawatt-hours — a Monthly Performance report of 2,989,497 kWh and a
  59% CUF for two solar Plants. PR was never computed at all: the irradiation
  passed to it was a literal zero.

PDF rendering uses WeasyPrint, which is an optional extra (`pip install -e
".[reports]"`) because it needs system pango/cairo. When it is absent the XLSX is
still produced and the PDF is reported as unavailable, rather than the whole run
failing: a Report that arrives in one format beats a Report that does not arrive.
"""

from __future__ import annotations

import html
import io
import math
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from typing import Any
from zoneinfo import ZoneInfo

import structlog
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from solarcms.domain.assumptions import (
    PLANT_ENERGY_COUNTER_PRECEDENCE,
    PLANT_IRRADIATION_SOURCE,
    TAG_SPECS,
)
from solarcms.domain.counters import (
    DeviceSeries,
    PlantEnergy,
    PlantIrradiation,
    bucket_steps,
    plant_energy,
    plant_irradiation,
)
from solarcms.domain.formulas import (
    FormulaResult,
    co2_avoided_kg,
    cuf,
    performance_ratio,
    specific_yield,
)
from solarcms.domain.tiering import TIERS, Tier, select_tier
from solarcms.services.energy import read_counter_series, split_by_pair

log = structlog.get_logger(__name__)

# Device Types whose Readings may be used for a Financial Report (I-11).
SETTLEMENT_DEVICE_TYPES = ("ABT_METER",)

# The client-level figures, in the order they are shown, with how each reads.
KPI_ROWS: tuple[tuple[str, str, str], ...] = (
    # key, label, unit
    ("energy_kwh", "Energy exported", "kWh"),
    ("specific_yield", "Specific yield", "kWh/kWp"),
    # The value already carries its % sign; a Unit column saying "%" as well
    # reads as a second, different figure.
    ("cuf", "CUF", ""),
    ("performance_ratio", "Performance ratio", ""),
    ("co2_avoided_kg", "CO₂ avoided", "kg"),
)
RATIO_KPIS = frozenset({"cuf", "performance_ratio"})

ANOMALY_WORDING = {
    "backwards": "went backwards (a reset, rollover or replaced meter)",
    "implausible_jump": "jumped further than the Plant could produce",
}


class FinancialSourceUnavailable(RuntimeError):
    """A Financial Report was requested but no ABT Meter Reading is available."""


@dataclass
class PlantReport:
    plant_id: int
    code: str
    name: str
    timezone: str
    dc_kwp: float
    ac_kw: float
    energy: PlantEnergy
    irradiation: PlantIrradiation
    kpis: dict[str, FormulaResult]
    # Local period label (ISO, sortable) -> kWh. Absent label = no readings.
    breakdown: dict[str, float] = field(default_factory=dict)


@dataclass
class ReportData:
    definition_code: str
    title: str
    period_start: datetime
    period_end: datetime
    tier: Tier
    client_name: str = ""
    is_financial: bool = False
    # "hour", "day", or "period" (no breakdown).
    granularity: str = "period"
    # The one timezone every Plant shares, or None when they differ (then UTC).
    timezone: str | None = None
    plants: list[PlantReport] = field(default_factory=list)
    # Per-meter detail, one row per Device that answered. `len(rows)` is what
    # the scheduler records as the run's row count.
    rows: list[dict[str, Any]] = field(default_factory=list)
    anomalies: list[dict[str, Any]] = field(default_factory=list)
    # Client totals: {"value", "variant", "undefined_reason"} per figure, and
    # `energy_kwh` as the same shape.
    kpis: dict[str, dict[str, Any]] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)


def _render(result: FormulaResult) -> dict[str, Any]:
    return {"value": result.value, "variant": result.variant,
            "undefined_reason": result.undefined_reason}


def _label_fn(granularity: str, tz: ZoneInfo) -> Callable[[datetime], str]:
    """A step's local hour or day — the Plant's own clock, never UTC, or an
    Asia/Kolkata Plant's morning lands on the previous day."""
    fmt = "%Y-%m-%d %H:00" if granularity == "hour" else "%Y-%m-%d"

    def label(at: datetime) -> str:
        return at.astimezone(tz).strftime(fmt)

    return label


def _all_labels(start: datetime, end: datetime, granularity: str, tz: ZoneInfo) -> list[str]:
    """Every local hour or day in the period, so a gap is shown as a gap."""
    step = timedelta(hours=1) if granularity == "hour" else timedelta(days=1)
    fmt = "%Y-%m-%d %H:00" if granularity == "hour" else "%Y-%m-%d"
    local = start.astimezone(tz)
    cursor = (local.replace(minute=0, second=0, microsecond=0) if granularity == "hour"
              else local.replace(hour=0, minute=0, second=0, microsecond=0))
    labels: list[str] = []
    end_local = end.astimezone(tz)
    while cursor < end_local and len(labels) < 20_000:
        labels.append(cursor.strftime(fmt))
        cursor = (cursor + step).astimezone(tz)
    return labels


@dataclass(frozen=True)
class PlantRecord:
    id: int
    code: str
    name: str
    timezone: str
    dc_capacity_kwp: float | None
    ac_capacity_kw: float | None
    grid_factor: float | None


async def gather(
    session: AsyncSession, *, definition_id: int, period_start: datetime,
    period_end: datetime, client_id: int,
) -> ReportData:
    """Collect the data one Report needs, from the correct tier.

    The reading is here; the arithmetic is `assemble`, which needs no database
    and is what the rendering tests drive.
    """
    definition = (await session.execute(text("""
        SELECT code, name, query_spec, is_financial
          FROM report_definitions WHERE id = :id
    """), {"id": definition_id})).first()
    if definition is None:
        raise ValueError(f"report definition {definition_id} not found")

    # Hourly at the finest: the scheduler role may read `agg_1h_v` and `agg_1d_v`
    # and nothing finer, and hourly is all a Report's breakdown needs.
    tier = select_tier(period_start, period_end, datetime.now(UTC), finest=Tier.AGG_1H)
    spec = definition.query_spec or {}
    precedence = _precedence(spec, is_financial=bool(definition.is_financial))

    client_name = (await session.execute(
        text("SELECT name FROM clients WHERE id = :id"), {"id": client_id})).scalar() or ""

    # ⚠ Explicit, and load-bearing. The scheduler runs with platform privileges
    # because it serves every Client, so neither RLS nor the barrier views scope
    # these rows. Without `client_id` a Report would include another Client's
    # Plants, and a Financial Report would invoice against another Client's meter.
    plants = [
        PlantRecord(
            id=row.id, code=row.code, name=row.name, timezone=row.timezone or "UTC",
            dc_capacity_kwp=float(row.dc_capacity_kwp) if row.dc_capacity_kwp else None,
            ac_capacity_kw=float(row.ac_capacity_kw) if row.ac_capacity_kw else None,
            grid_factor=float(row.grid_factor) if row.grid_factor is not None else None,
        )
        for row in (await session.execute(text("""
            SELECT p.id, p.code, p.name, p.timezone, p.dc_capacity_kwp, p.ac_capacity_kw,
                   r.grid_emission_factor_kg_per_kwh AS grid_factor
              FROM plants p LEFT JOIN regions r ON r.id = p.region_id
             WHERE p.client_id = :client_id AND p.status = 'active'
             ORDER BY p.code
        """), {"client_id": client_id})).all()
    ]

    series_by_plant = await read_counter_series(
        session, tier=tier, pairs=[*precedence, PLANT_IRRADIATION_SOURCE],
        start=period_start, end=period_end, plant_ids=[p.id for p in plants],
        client_id=client_id,
    ) if plants else {}

    return assemble(
        code=definition.code, title=definition.name, query_spec=spec,
        is_financial=bool(definition.is_financial), client_name=client_name,
        period_start=period_start, period_end=period_end, tier=tier,
        plants=plants, series_by_plant=series_by_plant,
    )


def _precedence(spec: dict[str, Any], *, is_financial: bool) -> tuple[tuple[str, str], ...]:
    if is_financial:
        # I-11 enforced by the precedence itself: one entry, no fallback.
        tag = str(spec.get("energy_tag", "ENERGY_EXPORT_TOTAL"))
        return tuple((type_code, tag) for type_code in SETTLEMENT_DEVICE_TYPES)
    return PLANT_ENERGY_COUNTER_PRECEDENCE


def assemble(
    *, code: str, title: str, query_spec: dict[str, Any], is_financial: bool,
    client_name: str, period_start: datetime, period_end: datetime, tier: Tier,
    plants: Sequence[PlantRecord], series_by_plant: dict[int, list[DeviceSeries]],
) -> ReportData:
    """Everything a Report says, from Plant records and their counter series."""
    resolution = next(s.resolution for s in TIERS if s.tier == tier)
    precedence = _precedence(query_spec, is_financial=is_financial)

    granularity = str(query_spec.get("granularity") or "period")
    notes: list[str] = []
    if granularity == "hour" and resolution > timedelta(hours=1):
        granularity = "day"
        notes.append(f"An hourly breakdown was asked for, but a period this long is read "
                     f"from the {tier.value} tier, so it is given per day.")
    if is_financial:
        notes.append("Financial Report: computed exclusively from ABT Meter Readings (I-11). "
                     "MFM Readings are excluded regardless of availability.")
    else:
        notes.append("Energy is read from one meter per Plant: the ABT Meter where it "
                     "reports, otherwise the MFM, otherwise the Inverters summed. That order "
                     "is provisional until the client says which register is commercially "
                     "binding (OPEN-14).")

    hours = max((period_end - period_start).total_seconds() / 3600.0, 1.0)
    data = ReportData(
        definition_code=code, title=title,
        period_start=period_start, period_end=period_end, tier=tier,
        client_name=client_name, is_financial=is_financial, granularity=granularity,
        timezone=plants[0].timezone if plants and len({p.timezone for p in plants}) == 1
        else None,
    )

    for plant in plants:
        stations, meters = split_by_pair(series_by_plant.get(plant.id, []),
                                         PLANT_IRRADIATION_SOURCE)
        dc = plant.dc_capacity_kwp or 0.0
        ac = plant.ac_capacity_kw or 0.0
        energy = plant_energy(meters, precedence, plant_ac_capacity_kw=ac or None)
        irradiation = plant_irradiation(stations)

        if energy.value is None:
            reason = energy.undefined_reason
            kpis = {
                "specific_yield": FormulaResult(None, "specific_yield", reason),
                "cuf": FormulaResult(None, cuf(0.0, 0.0, 0.0).variant, reason),
                "performance_ratio": FormulaResult(
                    None, performance_ratio(0.0, 0.0, 0.0).variant, reason),
                "co2_avoided_kg": FormulaResult(
                    None, co2_avoided_kg(0.0, plant.grid_factor).variant, reason),
            }
        else:
            kpis = {
                "specific_yield": specific_yield(energy.value, dc),
                "cuf": cuf(energy.value, ac, hours),
                "performance_ratio": performance_ratio(
                    energy.value, (irradiation.value or 0.0) * 1000.0, dc),
                "co2_avoided_kg": co2_avoided_kg(energy.value, plant.grid_factor),
            }
        tz = ZoneInfo(plant.timezone)
        report = PlantReport(
            plant_id=plant.id, code=plant.code, name=plant.name,
            timezone=plant.timezone, dc_kwp=dc, ac_kw=ac,
            energy=energy, irradiation=irradiation, kpis=kpis,
            breakdown=(bucket_steps(energy.steps, _label_fn(granularity, tz))
                       if granularity in ("hour", "day") else {}),
        )
        data.plants.append(report)

        for device in energy.devices:
            first, last = device.integral.first, device.integral.last
            data.rows.append({
                "plant_code": plant.code, "device_code": device.series.device_code,
                "device_type_code": device.series.device_type_code,
                "tag_code": device.series.tag_code,
                "unit": TAG_SPECS[device.series.tag_code].unit
                if device.series.tag_code in TAG_SPECS else "",
                "first_at": first.at if first else None,
                "first_value": first.value if first else None,
                "last_at": last.at if last else None,
                "last_value": last.value if last else None,
                "energy_kwh": device.integral.total,
                "samples": device.integral.samples,
                "refused": len(device.integral.anomalies),
                "timezone": report.timezone,
            })
        for series, anomaly in energy.anomalies:
            data.anomalies.append({
                "plant_code": plant.code, "device_code": series.device_code,
                "tag_code": series.tag_code, "kind": anomaly.kind,
                "previous_at": anomaly.previous_at, "at": anomaly.at,
                "from_value": anomaly.from_value, "to_value": anomaly.to_value,
                "limit": anomaly.limit, "timezone": report.timezone,
            })

    answered = [p for p in data.plants if p.energy.value is not None]
    if is_financial and not answered:
        raise FinancialSourceUnavailable(
            "no ABT Meter reported its export counter at least twice in this period; "
            "a Financial Report cannot be produced from an MFM (I-11)"
        )

    # ── Client totals ───────────────────────────────────────────────────────
    # Sums over the Plants that answered; a Plant with nothing to read is left
    # out of both sides rather than counted as zero.
    total = sum(p.energy.value or 0.0 for p in answered)
    dc_total = sum(p.dc_kwp for p in answered)
    ac_total = sum(p.ac_kw for p in answered)
    none_reason = "no Plant had an energy counter that reported twice in the period"
    data.kpis["energy_kwh"] = (
        {"value": total, "variant": "counter_steps", "undefined_reason": None}
        if answered else {"value": None, "variant": "counter_steps",
                          "undefined_reason": none_reason})
    if answered:
        data.kpis["specific_yield"] = _render(specific_yield(total, dc_total))
        data.kpis["cuf"] = _render(cuf(total, ac_total, hours))
        # PR pooled, not averaged: total energy over total reference yield,
        # across the Plants whose PR is defined (they had sun and a station).
        with_pr = [p for p in answered if p.kpis["performance_ratio"].value is not None]
        if with_pr:
            dc_pr = sum(p.dc_kwp for p in with_pr)
            weighted_irr = sum(
                p.dc_kwp * (p.irradiation.value or 0.0) * 1000.0 for p in with_pr) / dc_pr
            data.kpis["performance_ratio"] = _render(performance_ratio(
                sum(p.energy.value or 0.0 for p in with_pr), weighted_irr, dc_pr))
        else:
            data.kpis["performance_ratio"] = _render(answered[0].kpis["performance_ratio"])
        co2_values = [p.kpis["co2_avoided_kg"] for p in answered]
        variants = {c.variant for c in co2_values}
        data.kpis["co2_avoided_kg"] = {
            "value": sum(c.value or 0.0 for c in co2_values),
            "variant": variants.pop() if len(variants) == 1 else "co2_per_region",
            "undefined_reason": None,
        }
    else:
        for key, _label, _unit in KPI_ROWS[1:]:
            data.kpis[key] = {"value": None, "variant": "", "undefined_reason": none_reason}

    if data.anomalies:
        notes.append(
            f"{len(data.anomalies)} counter reading(s) were refused as generation — see "
            "'Refused readings'. The energy across each is not known, so the totals "
            "above are short by an amount this Report cannot state.")
    else:
        notes.append("No energy counter went backwards or jumped in this period.")
    notes.append("All KPI formulas are provisional pending OPEN-16; each figure records "
                 "the formula used so it can be recomputed when the client supplies theirs.")
    notes.append(f"Read from the {tier.value} tier, never raw readings.")
    data.notes = notes
    return data


# ── Presentation ─────────────────────────────────────────────────────────────

def _period_label(data: ReportData) -> str:
    """The period in the Plants' own time, as dates when it is whole days.

    ⚠ It used to print UTC clock times, so a month in Asia/Kolkata read
    "22-08-2026 18:30 to 22-09-2026 18:29" — correct, and a riddle.
    """
    tz = ZoneInfo(data.timezone or "UTC")
    start = data.period_start.astimezone(tz)
    end = data.period_end.astimezone(tz)
    whole_days = start.hour == 0 and start.minute == 0 and (
        (end.hour == 23 and end.minute == 59) or (end.hour == 0 and end.minute == 0))
    if whole_days:
        last_day = end if end.hour == 23 else end - timedelta(days=1)
        text_ = f"{start:%d-%m-%Y} to {last_day:%d-%m-%Y}"
    else:
        text_ = f"{start:%d-%m-%Y %H:%M} to {end:%d-%m-%Y %H:%M}"
    return f"{text_} ({data.timezone or 'UTC'})"


def _display_label(label: str) -> str:
    """An ISO period label as the tender's DD-MM-YYYY (§28), keeping any hour."""
    day = datetime.strptime(label[:10], "%Y-%m-%d").strftime("%d-%m-%Y")
    return f"{day} {label[11:]}" if len(label) > 10 else day


def _local_naive(at: datetime | None, timezone: str) -> datetime | None:
    # openpyxl cannot write a timezone-aware datetime; the zone is stated in
    # the column header instead.
    return at.astimezone(ZoneInfo(timezone)).replace(tzinfo=None) if at else None


def _source_label(energy: PlantEnergy) -> str:
    if energy.device_type_code is None:
        return energy.undefined_reason or "—"
    kind = {"ABT_METER": "ABT Meter", "MFM": "MFM", "INVERTER": "Inverters"}.get(
        energy.device_type_code, energy.device_type_code)
    if energy.device_type_code == "INVERTER":
        return f"{kind} (sum of {len(energy.devices)})"
    return kind


HEADER_FILL = PatternFill("solid", fgColor="E8ECF1")
BOLD = Font(bold=True)
MUTED = Font(color="6B7280")


def _header(sheet: Worksheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(sheet: Worksheet, widths: Sequence[float]) -> None:
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width


def _figure_cell(sheet: Worksheet, row: int, column: int, value: float | None,
                 number_format: str) -> None:
    """A number, or a grey dash — never 0 for "not known"."""
    if value is None:
        cell = sheet.cell(row=row, column=column, value="—")
        cell.font = MUTED
        cell.alignment = Alignment(horizontal="right")
        return
    cell = sheet.cell(row=row, column=column, value=value)
    cell.number_format = number_format


def _format_for(key: str) -> str:
    if key in RATIO_KPIS:
        return "0.00%"
    if key == "specific_yield":
        return "#,##0.00"
    if key == "co2_avoided_kg":
        return "#,##0"
    return "#,##0.0"


def render_xlsx(data: ReportData) -> bytes:
    """Summary, the breakdown, the meters used, and every refused reading."""
    book = Workbook()
    summary = book.active
    assert summary is not None
    summary.title = "Summary"
    _widths(summary, [30, 18, 16, 28, 18, 18, 30, 16])

    summary["A1"] = data.title
    summary["A1"].font = Font(size=16, bold=True)
    summary["A2"] = " · ".join(x for x in (data.client_name, _period_label(data)) if x)
    summary["A2"].font = MUTED

    row = 4
    _header(summary, row, ["Figure", "Value", "Unit", "Formula", "Note"])
    summary.merge_cells(start_row=row, start_column=5, end_row=row, end_column=8)
    for key, label, unit in KPI_ROWS:
        row += 1
        figure = data.kpis.get(key, {})
        summary.cell(row=row, column=1, value=label)
        _figure_cell(summary, row, 2, figure.get("value"), _format_for(key))
        summary.cell(row=row, column=3, value=unit).font = MUTED
        # The formula travels with the figure: a number that leaves the system
        # without its provenance cannot be reconciled later (OPEN-16).
        summary.cell(row=row, column=4, value=figure.get("variant") or "").font = MUTED
        if figure.get("undefined_reason"):
            summary.cell(row=row, column=5, value=figure["undefined_reason"]).font = MUTED
            summary.merge_cells(start_row=row, start_column=5, end_row=row, end_column=8)

    row += 2
    summary.cell(row=row, column=1, value="By Plant").font = Font(size=12, bold=True)
    row += 1
    _header(summary, row, ["Plant", "Energy (kWh)", "Specific yield (kWh/kWp)", "CUF",
                           "Performance ratio", "CO₂ avoided (kg)", "Energy source",
                           "Readings refused"])
    summary.row_dimensions[row].height = 30
    for plant in data.plants:
        row += 1
        summary.cell(row=row, column=1, value=f"{plant.name} ({plant.code})")
        _figure_cell(summary, row, 2, plant.energy.value, "#,##0.0")
        _figure_cell(summary, row, 3, plant.kpis["specific_yield"].value, "#,##0.00")
        _figure_cell(summary, row, 4, plant.kpis["cuf"].value, "0.00%")
        _figure_cell(summary, row, 5, plant.kpis["performance_ratio"].value, "0.00%")
        _figure_cell(summary, row, 6, plant.kpis["co2_avoided_kg"].value, "#,##0")
        summary.cell(row=row, column=7, value=_source_label(plant.energy))
        summary.cell(row=row, column=8, value=len(plant.energy.anomalies))
    if not data.plants:
        row += 1
        summary.cell(row=row, column=1, value="This Client has no active Plants.").font = MUTED

    row += 2
    for note in data.notes:
        cell = summary.cell(row=row, column=1, value=note)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.font = MUTED
        # Merged across the table's width, and tall enough for the text: a note
        # wrapped inside one narrow column reads one word per line.
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        summary.row_dimensions[row].height = 15 * max(1, math.ceil(len(note) / 140))
        row += 1

    # ── The breakdown ───────────────────────────────────────────────────────
    if data.granularity in ("hour", "day"):
        sheet = book.create_sheet("Generation")
        hourly = data.granularity == "hour"
        zone = data.timezone or "each Plant's own time"
        labels = ["Plant", f"Date ({zone})"] + ([f"Hour ({zone})"] if hourly else []) \
            + ["Energy (kWh)"]
        _header(sheet, 1, labels)
        _widths(sheet, [30, 16] + ([14] if hourly else []) + [16])
        line = 1
        for plant in data.plants:
            tz = ZoneInfo(plant.timezone)
            for label in _all_labels(data.period_start, data.period_end, data.granularity, tz):
                line += 1
                day = datetime.strptime(label[:10], "%Y-%m-%d")
                sheet.cell(row=line, column=1, value=plant.code)
                sheet.cell(row=line, column=2, value=day).number_format = "dd-mm-yyyy"
                column = 3
                if hourly:
                    sheet.cell(row=line, column=3, value=label[11:])
                    column = 4
                # No readings in that hour or day is a dash, never a zero.
                _figure_cell(sheet, line, column, plant.breakdown.get(label), "#,##0.0")
        sheet.freeze_panes = "A2"

    # ── The meters that answered ────────────────────────────────────────────
    detail = book.create_sheet("Data")
    _header(detail, 1, ["Plant", "Device", "Type", "Tag", "Unit", "First reading at",
                        "First reading", "Last reading at", "Last reading",
                        "Energy counted", "Readings", "Refused"])
    _widths(detail, [14, 16, 12, 22, 8, 18, 16, 18, 16, 16, 10, 10])
    for line, record in enumerate(data.rows, start=2):
        detail.cell(row=line, column=1, value=record["plant_code"])
        detail.cell(row=line, column=2, value=record["device_code"])
        detail.cell(row=line, column=3, value=record["device_type_code"])
        detail.cell(row=line, column=4, value=record["tag_code"])
        detail.cell(row=line, column=5, value=record["unit"])
        for column, key in ((6, "first_at"), (8, "last_at")):
            detail.cell(row=line, column=column,
                        value=_local_naive(record[key], record["timezone"])
                        ).number_format = "dd-mm-yyyy hh:mm"
        for column, key in ((7, "first_value"), (9, "last_value"), (10, "energy_kwh")):
            detail.cell(row=line, column=column, value=record[key]).number_format = "#,##0.0"
        detail.cell(row=line, column=11, value=record["samples"])
        detail.cell(row=line, column=12, value=record["refused"])
    detail.freeze_panes = "A2"

    # ── Every step refused as generation ────────────────────────────────────
    if data.anomalies:
        refused = book.create_sheet("Refused readings")
        _header(refused, 1, ["Plant", "Device", "Tag", "From (local)", "From value",
                             "To (local)", "To value", "What happened",
                             "Most it could have been (kWh)"])
        _widths(refused, [14, 16, 22, 18, 16, 18, 16, 48, 16])
        for line, item in enumerate(data.anomalies, start=2):
            refused.cell(row=line, column=1, value=item["plant_code"])
            refused.cell(row=line, column=2, value=item["device_code"])
            refused.cell(row=line, column=3, value=item["tag_code"])
            refused.cell(row=line, column=4, value=_local_naive(
                item["previous_at"], item["timezone"])).number_format = "dd-mm-yyyy hh:mm"
            refused.cell(row=line, column=5, value=item["from_value"]).number_format = "#,##0.0"
            refused.cell(row=line, column=6, value=_local_naive(
                item["at"], item["timezone"])).number_format = "dd-mm-yyyy hh:mm"
            refused.cell(row=line, column=7, value=item["to_value"]).number_format = "#,##0.0"
            refused.cell(row=line, column=8, value=ANOMALY_WORDING[item["kind"]])
            _figure_cell(refused, line, 9, item["limit"], "#,##0.0")
        refused.freeze_panes = "A2"

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _fmt(value: float | None, key: str) -> str:
    if value is None:
        return "—"
    if key in RATIO_KPIS:
        return f"{value * 100:.2f}%"
    if key == "specific_yield":
        return f"{value:,.2f}"
    if key == "co2_avoided_kg":
        return f"{value:,.0f}"
    return f"{value:,.1f}"


def render_html(data: ReportData) -> str:
    """The PDF source. Kept separate so it is renderable without WeasyPrint."""
    esc = html.escape

    def kpi_row(key: str, label: str, unit: str) -> str:
        figure = data.kpis.get(key, {})
        note = figure.get("undefined_reason") or figure.get("variant") or ""
        return (f"<tr><td>{esc(label)}</td><td class='n'>{_fmt(figure.get('value'), key)}</td>"
                f"<td>{esc(unit)}</td><td class='m'>{esc(note)}</td></tr>")

    kpis = "".join(kpi_row(key, label, unit) for key, label, unit in KPI_ROWS)
    plants = "".join(
        f"<tr><td>{esc(p.name)} ({esc(p.code)})</td>"
        f"<td class='n'>{_fmt(p.energy.value, 'energy_kwh')}</td>"
        f"<td class='n'>{_fmt(p.kpis['specific_yield'].value, 'specific_yield')}</td>"
        f"<td class='n'>{_fmt(p.kpis['cuf'].value, 'cuf')}</td>"
        f"<td class='n'>{_fmt(p.kpis['performance_ratio'].value, 'performance_ratio')}</td>"
        f"<td>{esc(_source_label(p.energy))}</td>"
        f"<td class='n'>{len(p.energy.anomalies)}</td></tr>"
        for p in data.plants
    )

    breakdown = ""
    if data.granularity in ("hour", "day"):
        sections = []
        for plant in data.plants:
            labels = _all_labels(data.period_start, data.period_end, data.granularity,
                                 ZoneInfo(plant.timezone))
            # A month of hours is 744 rows; the PDF keeps to days and points at
            # the workbook for the rest.
            if len(labels) > 62:
                sections.append(f"<p class='m'>{esc(plant.code)}: {len(labels)} "
                                f"{data.granularity}s — see the XLSX for the breakdown.</p>")
                continue
            rows_html = "".join(
                f"<tr><td>{esc(_display_label(label))}</td>"
                f"<td class='n'>{_fmt(plant.breakdown.get(label), 'energy_kwh')}</td></tr>"
                for label in labels)
            sections.append(f"<h3>{esc(plant.name)} ({esc(plant.code)})</h3><table>"
                            f"<tr><th>{'Hour' if data.granularity == 'hour' else 'Date'}</th>"
                            f"<th>Energy (kWh)</th></tr>{rows_html}</table>")
        breakdown = "<h2>Generation</h2>" + "".join(sections)

    refused = ""
    if data.anomalies:
        refused_rows = "".join(
            f"<tr><td>{esc(a['plant_code'])}</td><td>{esc(a['device_code'])}</td>"
            f"<td>{a['previous_at'].astimezone(ZoneInfo(a['timezone'])):%d-%m-%Y %H:%M}</td>"
            f"<td class='n'>{a['from_value']:,.1f}</td>"
            f"<td>{a['at'].astimezone(ZoneInfo(a['timezone'])):%d-%m-%Y %H:%M}</td>"
            f"<td class='n'>{a['to_value']:,.1f}</td>"
            f"<td>{esc(ANOMALY_WORDING[a['kind']])}</td></tr>"
            for a in data.anomalies)
        refused = ("<h2>Refused readings</h2><table><tr><th>Plant</th><th>Device</th>"
                   "<th>From</th><th>Value</th><th>To</th><th>Value</th><th>What happened</th>"
                   f"</tr>{refused_rows}</table>")

    notes = "".join(f"<li>{esc(n)}</li>" for n in data.notes)
    # Dates as DD-MM-YYYY per tender §28.
    return f"""<!doctype html><html><head><meta charset="utf-8">
<style>
 body {{ font-family: sans-serif; font-size: 11px; margin: 24px; color: #111; }}
 h1 {{ font-size: 18px; margin-bottom: 2px; }}
 h2 {{ font-size: 13px; margin: 18px 0 6px; }}
 h3 {{ font-size: 11px; margin: 10px 0 4px; }}
 .period, .m {{ color: #555; }}
 .period {{ margin-bottom: 16px; }}
 table {{ border-collapse: collapse; width: 100%; margin-bottom: 12px; }}
 th, td {{ border: 1px solid #d4d8de; padding: 4px 6px; text-align: left; }}
 th {{ background: #eef1f5; }}
 td.n {{ text-align: right; font-variant-numeric: tabular-nums; }}
 ul {{ color: #555; font-size: 10px; padding-left: 16px; }}
</style></head><body>
<h1>{esc(data.title)}</h1>
<div class="period">{esc(" · ".join(x for x in (data.client_name, _period_label(data)) if x))}
</div>
<h2>Key figures</h2>
<table><tr><th>Figure</th><th>Value</th><th>Unit</th><th>Formula / note</th></tr>{kpis}</table>
<h2>By Plant</h2>
<table><tr><th>Plant</th><th>Energy (kWh)</th><th>Specific yield</th><th>CUF</th><th>PR</th>
<th>Energy source</th><th>Refused</th></tr>{plants}</table>
{breakdown}
{refused}
<ul>{notes}</ul>
</body></html>"""


def render_pdf(data: ReportData) -> bytes | None:
    """PDF, or None when WeasyPrint is unavailable.

    Returning None rather than raising: the XLSX has already been produced by the
    time this runs, and failing the whole run over an optional system dependency
    would withhold a Report that is otherwise complete.
    """
    try:
        from weasyprint import HTML
    except ImportError:
        log.warning("weasyprint unavailable; PDF skipped",
                    hint='install with: pip install -e ".[reports]"')
        return None
    return bytes(HTML(string=render_html(data)).write_pdf())
