"""The Reports screen's tables: one Plant, whole local days, nothing invented.

Driven through the pure `build_*` functions, which need no database. The Plant
is in Kolkata so that every day boundary is half an hour off a UTC hour — the
shape that put a whole day's energy on the wrong date when daily buckets were
used.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime, timedelta
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from solarcms.domain.counters import CounterSample, DeviceSeries
from solarcms.domain.formulas import FormulaResult
from solarcms.domain.periods import report_window
from solarcms.domain.tiering import Tier
from solarcms.services.report_tables import (
    MAX_ALARM_ROWS,
    AlarmRecord,
    DeviceRecord,
    PlantInfo,
    Stat,
    build_alarm,
    build_daily_plant,
    build_inverter,
    build_monthly_plant,
    build_weather,
    device_availability,
    pick_source,
    render_csv,
    render_html,
    render_xlsx,
    split_steps,
    to_payload,
)

IST = ZoneInfo("Asia/Kolkata")
PLANT = PlantInfo(id=1, client_id=7, code="SF_NORTH", name="Sunfield North",
                  timezone="Asia/Kolkata", dc_kwp=1_200.0, ac_kw=1_000.0)
# 16:18 in Kolkata on 28 Sep 2026.
NOW = datetime(2026, 9, 28, 10, 48, tzinfo=UTC)
TIERS = {"counter_tier": Tier.AGG_1H, "stats_tier": Tier.AGG_15M}


def _local(day: int, hour: int, minute: int = 0, month: int = 9) -> datetime:
    return datetime(2026, month, day, hour, minute, tzinfo=IST).astimezone(UTC)


def _register(
    type_code: str, tag: str, readings: list[tuple[datetime, float]], *,
    device_id: int = 10, code: str = "MFM_1", rated: float | None = None,
) -> DeviceSeries:
    return DeviceSeries(
        device_id=device_id, device_code=code, device_type_code=type_code, tag_code=tag,
        samples=tuple(CounterSample(at, value) for at, value in readings),
        rated_capacity_kw=rated)


def _meter(*days: tuple[int, float]) -> DeviceSeries:
    """An MFM export register climbing `kwh` over each local day's daylight."""
    readings: list[tuple[datetime, float]] = []
    total = 50_000.0
    for day, kwh in days:
        readings.append((_local(day, 6), total))
        total += kwh
        readings.append((_local(day, 18), total))
    return _register("MFM", "ENERGY_EXPORT_TOTAL", readings)


class TestDailyPlant:
    WINDOW = report_window("last_7_days", NOW, IST)

    def test_one_row_per_local_day_with_that_days_energy(self) -> None:
        table = build_daily_plant(PLANT, self.WINDOW, [_meter((24, 4_000.0), (25, 5_000.0))],
                                  {}, **TIERS)
        assert [row["date"] for row in table.rows] == [
            f"2026-09-{d}" for d in range(22, 29)]
        by_date = {row["date"]: row for row in table.rows}
        assert by_date["2026-09-24"]["energy_kwh"] == 4_000.0
        assert by_date["2026-09-25"]["energy_kwh"] == 5_000.0

    def test_a_day_with_nothing_to_read_is_none_never_zero(self) -> None:
        table = build_daily_plant(PLANT, self.WINDOW, [_meter((24, 4_000.0))], {}, **TIERS)
        by_date = {row["date"]: row for row in table.rows}
        assert by_date["2026-09-22"]["energy_kwh"] is None
        assert by_date["2026-09-22"]["avg_power_kw"] is None

    def test_a_meter_register_is_called_export_and_an_inverter_one_is_not(self) -> None:
        meter = build_daily_plant(PLANT, self.WINDOW, [_meter((24, 1.0))], {}, **TIERS)
        inverter = build_daily_plant(PLANT, self.WINDOW, [_register(
            "INVERTER", "ENERGY_TOTAL", [(_local(24, 6), 0.0), (_local(24, 18), 9.0)])],
            {}, **TIERS)
        heading = {c.key: c.label for c in meter.columns}["energy_kwh"]
        assert heading == "Export energy"
        assert {c.key: c.label for c in inverter.columns}["energy_kwh"] == "Energy generated"

    def test_avg_power_is_energy_over_the_days_hours(self) -> None:
        table = build_daily_plant(PLANT, self.WINDOW, [_meter((24, 4_800.0))], {}, **TIERS)
        by_date = {row["date"]: row for row in table.rows}
        assert by_date["2026-09-24"]["avg_power_kw"] == 200.0

    def test_today_divides_only_by_the_hours_so_far(self) -> None:
        window = report_window("today", NOW, IST)
        series = _register("MFM", "ENERGY_EXPORT_TOTAL",
                           [(_local(28, 6), 100.0), (_local(28, 16), 1_730.0)])
        row = build_daily_plant(PLANT, window, [series], {}, **TIERS).rows[0]
        # 00:00 to 16:18 is 16.3 hours.
        assert row["avg_power_kw"] == 1_630.0 / 16.3

    def test_peak_and_frequency_come_from_the_first_source_that_reported(self) -> None:
        day = date(2026, 9, 24)
        stats = {
            ("MFM", "AC_ACTIVE_POWER"): {day: Stat(810.0, 300.0, 830.0, 96)},
            ("INVERTER", "AC_ACTIVE_POWER"): {day: Stat(790.0, 290.0, 60.0, 96 * 12,
                                                        devices=12)},
            ("INVERTER", "FREQUENCY"): {day: Stat(50.1, 50.02, 50.1, 400)},
        }
        table = build_daily_plant(PLANT, self.WINDOW, [], stats, **TIERS)
        row = {r["date"]: r for r in table.rows}["2026-09-24"]
        assert row["peak_power_kw"] == 810.0
        assert row["avg_frequency_hz"] == 50.02
        assert any("the MFM" in note for note in table.notes)

    def test_the_inverter_count_in_the_note_is_read_from_the_data(self) -> None:
        day = date(2026, 9, 24)
        stats = {("INVERTER", "AC_ACTIVE_POWER"): {day: Stat(790.0, 290.0, 60.0, 900,
                                                             devices=17)}}
        table = build_daily_plant(PLANT, self.WINDOW, [], stats, **TIERS)
        assert any("the 17 Inverters summed" in note for note in table.notes)

    def test_import_without_a_meter_is_unknown_and_says_why(self) -> None:
        table = build_daily_plant(PLANT, self.WINDOW, [_meter((24, 1.0))], {}, **TIERS)
        assert all(row["import_kwh"] is None for row in table.rows)
        assert any("Import is not known" in note for note in table.notes)

    def test_flagged_buckets_are_counted_in_the_notes(self) -> None:
        day = date(2026, 9, 24)
        stats = {("MFM", "AC_ACTIVE_POWER"): {day: Stat(810.0, 300.0, 830.0, 96, flagged=3)}}
        table = build_daily_plant(PLANT, self.WINDOW, [], stats, **TIERS)
        assert any(note.startswith("3 bucket(s) held a flagged reading")
                   for note in table.notes)


class TestOutages:
    """A register step across an outage belongs to no single day."""

    def _gap(self) -> DeviceSeries:
        # Read through the 25th, silent on the 26th and 27th, back at 16:00 on
        # the 28th having climbed 277 kWh at some unknown point in between.
        return _register("MFM", "ENERGY_EXPORT_TOTAL", [
            (_local(25, 6), 1_000.0), (_local(25, 18), 5_000.0),
            (_local(28, 16), 5_277.0), (_local(28, 16, 15), 5_300.0)])

    def test_the_step_across_the_gap_is_in_no_row_and_is_stated(self) -> None:
        table = build_daily_plant(PLANT, report_window("last_7_days", NOW, IST),
                                  [self._gap()], {}, **TIERS)
        by_date = {row["date"]: row for row in table.rows}
        assert by_date["2026-09-25"]["energy_kwh"] == 4_000.0
        assert by_date["2026-09-28"]["energy_kwh"] == 23.0
        assert any(note.startswith("277.0 kWh of energy accrued") for note in table.notes)

    def test_a_day_reads_the_same_whichever_period_contains_it(self) -> None:
        # The regression: "today" never saw the step, "last 7 days" gave it
        # all to today, and the same day carried two totals.
        today_window = report_window("today", NOW, IST)
        in_today = [s for s in self._gap().samples if s.at >= today_window.start]
        today = build_daily_plant(PLANT, today_window, [_register(
            "MFM", "ENERGY_EXPORT_TOTAL", [(s.at, s.value) for s in in_today])], {}, **TIERS)
        week = build_daily_plant(PLANT, report_window("last_7_days", NOW, IST),
                                 [self._gap()], {}, **TIERS)
        assert today.rows[0]["energy_kwh"] == week.rows[-1]["energy_kwh"]

    def test_split_steps_keeps_steps_whose_readings_share_the_day(self) -> None:
        series = _meter((24, 100.0))
        from solarcms.domain.counters import integrate_counter

        steps = integrate_counter(series.samples, max_rate_per_hour=None).steps
        split = split_steps(steps, lambda at: at.astimezone(IST).date().isoformat())
        assert split.totals == {"2026-09-24": 100.0}
        assert split.unattributed == 0.0


    def test_a_step_ending_exactly_at_midnight_is_the_day_befores(self) -> None:
        # A quarter-hour bucket's last reading is stamped at the bucket's end,
        # so the day's final step ends on the stroke of midnight.
        series = _register("MFM", "ENERGY_EXPORT_TOTAL", [
            (_local(24, 23, 45), 100.0), (_local(25, 0), 104.0), (_local(25, 0, 15), 105.0)])
        table = build_daily_plant(PLANT, report_window("last_7_days", NOW, IST),
                                  [series], {}, **TIERS)
        by_date = {row["date"]: row for row in table.rows}
        assert by_date["2026-09-24"]["energy_kwh"] == 4.0
        assert by_date["2026-09-25"]["energy_kwh"] == 1.0
        assert not any("accrued while" in note for note in table.notes)


class TestCoverage:
    def test_a_day_seen_for_part_of_its_slots_says_so(self) -> None:
        day = date(2026, 9, 24)
        stats = {("MFM", "AC_ACTIVE_POWER"): {day: Stat(810.0, 300.0, 830.0, 48,
                                                        buckets=48)}}
        table = build_daily_plant(PLANT, report_window("last_7_days", NOW, IST), [], stats,
                                  **TIERS)
        by_date = {row["date"]: row for row in table.rows}
        # 48 of the day's 96 quarter-hours.
        assert by_date["2026-09-24"]["coverage"] == 50.0
        assert by_date["2026-09-23"]["coverage"] == 0.0

    def test_no_power_source_at_all_is_unknown_not_zero(self) -> None:
        table = build_daily_plant(PLANT, report_window("yesterday", NOW, IST), [], {},
                                  **TIERS)
        assert table.rows[0]["coverage"] is None


class TestEmptyColumns:
    def test_a_column_nothing_reported_is_explained_once(self) -> None:
        table = build_daily_plant(PLANT, report_window("yesterday", NOW, IST),
                                  [_meter((27, 10.0))], {}, **TIERS)
        empty = [n for n in table.notes if n.startswith("Nothing reported")]
        assert empty == ["Nothing reported Peak power, Import energy, Avg frequency, "
                         "Coverage in the period; those columns are empty, not zero."]


class TestPickSource:
    def test_skips_a_pair_whose_days_hold_no_good_reading(self) -> None:
        day = date(2026, 9, 24)
        stats = {("MFM", "AC_ACTIVE_POWER"): {day: Stat(None, None, None, 0, flagged=4)},
                 ("INVERTER", "AC_ACTIVE_POWER"): {day: Stat(700.0, 200.0, 70.0, 10)}}
        pair, _days = pick_source(stats, (("MFM", "AC_ACTIVE_POWER"),
                                          ("INVERTER", "AC_ACTIVE_POWER")))
        assert pair == ("INVERTER", "AC_ACTIVE_POWER")

    def test_none_when_nothing_reported(self) -> None:
        assert pick_source({}, (("MFM", "AC_ACTIVE_POWER"),)) == (None, {})


class TestMonthlyPlant:
    def test_one_row_per_month_the_period_touches(self) -> None:
        window = report_window("last_30_days", NOW, IST)
        table = build_monthly_plant(PLANT, window, [], {}, **TIERS)
        assert [(row["month"], row["days"]) for row in table.rows] == [
            ("2026-08", 2), ("2026-09", 28)]

    def test_pr_is_a_percentage_and_an_impossible_one_is_flagged_not_clamped(self) -> None:
        window = report_window("custom", NOW, IST, date(2026, 9, 24), date(2026, 9, 24))
        station = _register("WMS", "GHI_CUMULATIVE",
                            [(_local(24, 6), 0.0), (_local(24, 18), 0.5)],
                            device_id=20, code="WMS_1")
        # 1,200 kWh from 1,200 kWp is 1 kWh/kWp over 0.5 kWh/m² of sun: PR 200%.
        table = build_monthly_plant(PLANT, window, [_meter((24, 1_200.0)), station], {},
                                    **TIERS)
        row = table.rows[0]
        assert row["performance_ratio"] == 200.0
        assert (0, "performance_ratio") in [(i, k) for i, k, _ in table.flags]

    def test_cuf_counts_only_the_hours_since_the_plant_was_first_heard(self) -> None:
        window = report_window("custom", NOW, IST, date(2026, 9, 1), date(2026, 9, 24))
        table = build_monthly_plant(PLANT, window, [_meter((24, 2_400.0))], {}, **TIERS)
        # First reading 06:00 on the 24th; the window ends at 00:00 on the 25th.
        assert table.rows[0]["cuf"] == 2_400.0 / (1_000.0 * 18) * 100.0


class TestInverter:
    WINDOW = report_window("yesterday", NOW, IST)

    def test_each_inverters_own_register_and_peaks(self) -> None:
        inverters = [DeviceRecord(1, "INV_01", "Inverter 1", _local(1, 0)),
                     DeviceRecord(2, "INV_02", "Inverter 2", _local(1, 0))]
        series = [_register("INVERTER", "ENERGY_TOTAL",
                            [(_local(27, 6), 1_000.0), (_local(27, 18), 1_420.0)],
                            device_id=1, code="INV_01", rated=100.0)]
        stats = {(1, "AC_ACTIVE_POWER"): Stat(98.0, 40.0, 98.0, 96)}
        uptime = {1: FormulaResult(0.975, "availability"),
                  2: FormulaResult(None, "availability", "never reported")}
        table = build_inverter(PLANT, self.WINDOW, inverters, series, stats, uptime, **TIERS)
        first, second = table.rows
        assert first["energy_kwh"] == 420.0
        assert first["peak_ac_kw"] == 98.0
        assert first["availability"] == 97.5
        # An Inverter that said nothing is a row of dashes, not of zeros.
        assert second["energy_kwh"] is None
        assert second["availability"] is None

    def test_availability_excludes_time_before_registration(self) -> None:
        registered = _local(27, 12)
        result = device_availability([(registered, "online")], registered, self.WINDOW)
        assert result.value == 1.0


class TestWeather:
    def test_irradiation_is_the_mean_of_the_stations_that_read_that_day(self) -> None:
        window = report_window("custom", NOW, IST, date(2026, 9, 24), date(2026, 9, 25))
        a = _register("WMS", "GHI_CUMULATIVE",
                      [(_local(24, 6), 0.0), (_local(24, 18), 5.0),
                       (_local(25, 6), 0.0), (_local(25, 18), 4.0)], device_id=1, code="W1")
        b = _register("WMS", "GHI_CUMULATIVE",
                      [(_local(24, 6), 0.0), (_local(24, 18), 6.0)], device_id=2, code="W2")
        table = build_weather(PLANT, window, [a, b], {}, **TIERS)
        assert table.rows[0]["ghi_kwh_m2"] == 5.5
        # W2 was down on the 25th: W1 alone, not W1 halved.
        assert table.rows[1]["ghi_kwh_m2"] == 4.0


class TestAlarm:
    WINDOW = report_window("today", NOW, IST)

    def _alarm(self, minute: int, resolved: bool) -> AlarmRecord:
        opened = _local(28, 9, minute)
        return AlarmRecord(opened, "high", "INV_03", "Communication lost", "communication",
                           "resolved" if resolved else "active", None,
                           opened + timedelta(minutes=45) if resolved else None, "silent")

    def test_times_are_the_plants_and_an_open_alarm_has_no_duration(self) -> None:
        table = build_alarm(PLANT, self.WINDOW, [self._alarm(0, True), self._alarm(5, False)])
        assert table.rows[0]["opened_at"] == "2026-09-28T09:00:00+05:30"
        assert table.rows[0]["duration_min"] == 45.0
        assert table.rows[1]["duration_min"] is None

    def test_a_long_list_is_cut_and_says_so(self) -> None:
        alarms = [self._alarm(0, True)] * (MAX_ALARM_ROWS + 1)
        table = build_alarm(PLANT, self.WINDOW, alarms)
        assert len(table.rows) == MAX_ALARM_ROWS
        assert table.truncated


class TestRendering:
    def _table(self):  # type: ignore[no-untyped-def]
        window = report_window("last_7_days", NOW, IST)
        return build_daily_plant(PLANT, window, [_meter((24, 4_800.0))], {}, **TIERS)

    def test_csv_headings_carry_units_and_unknown_is_empty(self) -> None:
        rows = list(csv.reader(io.StringIO(render_csv(self._table()))))
        assert rows[0][:4] == ["Date", "Avg power (kW)", "Peak power (kW)",
                               "Export energy (kWh)"]
        by_date = {row[0]: row for row in rows[1:]}
        assert by_date["24-09-2026"][3] == "4800.0"
        assert by_date["22-09-2026"][3] == ""

    def test_xlsx_writes_a_dash_for_unknown_never_zero(self) -> None:
        book = load_workbook(io.BytesIO(render_xlsx(self._table())))
        sheet = book.active
        assert sheet is not None
        assert sheet["A1"].value == "Daily Plant Report"
        # Header on row 4; the first day, the 22nd, has nothing to read.
        assert sheet.cell(row=5, column=4).value == "—"
        assert sheet.cell(row=7, column=4).value == 4_800.0

    def test_html_and_payload_name_the_plant_and_period(self) -> None:
        table = self._table()
        page = render_html(table)
        assert "Daily Plant Report" in page and "SF_NORTH" in page
        assert "22-09-2026 to 28-09-2026" in page
        payload = to_payload(table, NOW)
        assert payload["first_day"] == "2026-09-22"
        assert payload["columns"][0] == {"key": "date", "label": "Date", "kind": "date",
                                         "unit": "", "digits": 1}
