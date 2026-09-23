"""A Report says one figure per Plant, from one meter, and shows what it refused.

Driven through `reporting.assemble`, which needs no database. The fleet is the
shape that broke the old Report: a Plant with both a settlement meter and a
check meter (counted twice, before), and a meter whose counter restarted and
went backwards (counted as gigawatt-hours, before).
"""

from __future__ import annotations

import io
import math
from datetime import UTC, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from solarcms.domain.counters import CounterSample, DeviceSeries
from solarcms.domain.tiering import Tier
from solarcms.services.reporting import (
    FinancialSourceUnavailable,
    PlantRecord,
    ReportData,
    assemble,
    render_html,
    render_xlsx,
)

IST = ZoneInfo("Asia/Kolkata")
START = datetime(2026, 9, 20, tzinfo=IST).astimezone(UTC)
END = datetime(2026, 9, 22, 23, 59, tzinfo=IST).astimezone(UTC)
HOURS = [START + timedelta(hours=h + 1) for h in range(72)]


def _solar_kwh(at: datetime, capacity_kw: float) -> float:
    """One hour's generation on a clear day, in the Plant's own clock."""
    hour = at.astimezone(IST).hour
    return 0.0 if not 7 <= hour <= 18 else capacity_kw * 0.8 * math.sin((hour - 6) / 13 * math.pi)


def _counter(code: str, type_code: str, capacity_kw: float, base: float,
             bias: float = 1.0, *, jump_at: int | None = None,
             back_at: int | None = None) -> DeviceSeries:
    value, samples = base, []
    for index, at in enumerate(HOURS):
        value += _solar_kwh(at, capacity_kw) * bias
        if index == jump_at:
            value += 800_000.0          # a restarted counter, not generation
        if index == back_at:
            value -= 250_000.0          # and one that went backwards
        samples.append(CounterSample(at, value))
    return DeviceSeries(1, code, type_code, "ENERGY_EXPORT_TOTAL", tuple(samples))


def _station(code: str) -> DeviceSeries:
    value, samples = 0.0, []
    for at in HOURS:
        if at.astimezone(IST).hour == 0:
            value = 0.0                  # the daily register resets at midnight
        value += _solar_kwh(at, 1.0)     # ~kWh/m2 per hour at 1 kW/m2 peak
        samples.append(CounterSample(at, value))
    return DeviceSeries(9, code, "WMS", "GHI_CUMULATIVE", tuple(samples))


PLANTS = [
    PlantRecord(1, "SF_NORTH", "Sunfield North", "Asia/Kolkata", 5760.0, 4800.0, None),
    PlantRecord(2, "SF_SOUTH", "Sunfield South", "Asia/Kolkata", 2400.0, 2000.0, None),
]
SERIES = {
    1: [_counter("ABT_METER", "ABT_METER", 4800.0, 1_440_000.0, 0.977),
        _counter("MFM", "MFM", 4800.0, 1_230_000.0, 0.98),
        _station("WMS")],
    2: [_counter("MFM", "MFM", 2000.0, 1_850_000.0, 0.98, jump_at=30, back_at=50),
        _station("WMS")],
}


def _report(code: str = "monthly_performance", granularity: str = "day",
            financial: bool = False) -> ReportData:
    return assemble(
        code=code, title="Monthly Performance", is_financial=financial,
        query_spec={"energy_tag": "ENERGY_EXPORT_TOTAL", "granularity": granularity},
        client_name="Sunfield Energy", period_start=START, period_end=END,
        tier=Tier.AGG_1H, plants=PLANTS, series_by_plant=SERIES,
    )


def _expected(capacity: float, bias: float, skip: frozenset[int] = frozenset()) -> float:
    return sum(_solar_kwh(at, capacity) * bias for i, at in enumerate(HOURS[1:], 1)
               if i not in skip)


class TestFigures:
    def test_one_meter_per_plant_never_both(self) -> None:
        data = _report()
        north = next(p for p in data.plants if p.code == "SF_NORTH")
        assert north.energy.device_type_code == "ABT_METER"
        assert north.energy.value == pytest.approx(_expected(4800.0, 0.977))

    def test_a_restarted_counter_is_refused_and_listed(self) -> None:
        data = _report()
        south = next(p for p in data.plants if p.code == "SF_SOUTH")
        # The hours containing the jump and the drop are refused whole.
        assert south.energy.value == pytest.approx(_expected(2000.0, 0.98, frozenset({30, 50})))
        assert sorted(a["kind"] for a in data.anomalies) == ["backwards", "implausible_jump"]

    def test_the_total_is_the_two_plants_not_three_meters(self) -> None:
        data = _report()
        expected = _expected(4800.0, 0.977) + _expected(2000.0, 0.98, frozenset({30, 50}))
        assert data.kpis["energy_kwh"]["value"] == pytest.approx(expected)

    def test_pr_is_computed_from_the_weather_station(self) -> None:
        # It used to be passed an irradiation of literally zero.
        data = _report()
        pr = data.kpis["performance_ratio"]["value"]
        assert pr is not None and 0.5 < pr < 1.0

    def test_a_financial_report_uses_the_abt_meter_or_refuses(self) -> None:
        data = _report("monthly_settlement", financial=True)
        assert [p.energy.device_type_code for p in data.plants] == ["ABT_METER", None]
        with pytest.raises(FinancialSourceUnavailable, match="ABT Meter"):
            assemble(
                code="monthly_settlement", title="Settlement", is_financial=True,
                query_spec={}, client_name="", period_start=START, period_end=END,
                tier=Tier.AGG_1H, plants=PLANTS[1:], series_by_plant=SERIES)


class TestWorkbook:
    def test_it_reads_like_a_report(self) -> None:
        book = load_workbook(io.BytesIO(render_xlsx(_report())))
        summary = book["Summary"]
        assert summary["A2"].value == "Sunfield Energy · 20-09-2026 to 22-09-2026 (Asia/Kolkata)"
        assert summary.column_dimensions["A"].width >= 28, "labels were cut off at 8 chars"
        labels = [summary.cell(row=r, column=1).value for r in range(5, 10)]
        assert labels == ["Energy exported", "Specific yield", "CUF", "Performance ratio",
                          "CO₂ avoided"]
        cuf_cell = summary.cell(row=7, column=2)
        assert cuf_cell.number_format == "0.00%", "CUF printed as a bare fraction"
        # Notes span the table instead of wrapping one word per line in column A.
        assert any(str(r).startswith("A") and ":H" in str(r)
                   for r in summary.merged_cells.ranges)

    def test_every_local_day_is_listed(self) -> None:
        book = load_workbook(io.BytesIO(render_xlsx(_report())))
        rows = [r for r in book["Generation"].iter_rows(min_row=2, values_only=True)]
        north = [r for r in rows if r[0] == "SF_NORTH"]
        assert len(north) == 3
        assert all(isinstance(r[2], float) and r[2] > 0 for r in north)

    def test_refused_readings_have_their_own_sheet(self) -> None:
        book = load_workbook(io.BytesIO(render_xlsx(_report())))
        assert "Refused readings" in book.sheetnames
        assert book["Refused readings"].max_row == 3

    def test_an_hourly_report_lists_every_hour(self) -> None:
        book = load_workbook(io.BytesIO(render_xlsx(_report("daily_generation", "hour"))))
        rows = [r for r in book["Generation"].iter_rows(min_row=2, values_only=True)
                if r[0] == "SF_NORTH"]
        assert len(rows) == 72
        night = [r for r in rows if r[2] == "02:00"]
        assert night and all(r[3] == 0 for r in night), "a night hour made nothing, not '—'"

    def test_the_pdf_source_renders(self) -> None:
        page = render_html(_report())
        assert "Sunfield North (SF_NORTH)" in page and "Refused readings" in page
